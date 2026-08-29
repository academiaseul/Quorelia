# -*- coding: utf-8 -*-
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

AZ="0000FF"; NG="000000"; VD="008000"
AM=PatternFill("solid",fgColor="FFFF00")
CAB=PatternFill("solid",fgColor="0B1F2A")
GRIS=PatternFill("solid",fgColor="EEF1F3")
TEAL=PatternFill("solid",fgColor="D6F0EC")
ROJO=PatternFill("solid",fgColor="FBE0E0")
F="Arial"
hair=Side(style="thin",color="BBBBBB")
BOX=Border(left=hair,right=hair,top=hair,bottom=hair)
CLP='#,##0;(#,##0);-'; USD='"US$"#,##0;("US$"#,##0);-'; PCT='0.0%'; NUM='#,##0.0'

wb=Workbook()

def titulo(ws,txt,sub,ncol=7):
    ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=ncol)
    c=ws.cell(1,1,txt); c.font=Font(F,size=14,bold=True,color="FFFFFF"); c.fill=CAB
    c.alignment=Alignment(vertical="center",indent=1); ws.row_dimensions[1].height=30
    for j in range(2,ncol+1): ws.cell(1,j).fill=CAB
    ws.merge_cells(start_row=2,start_column=1,end_row=2,end_column=ncol)
    c=ws.cell(2,1,sub); c.font=Font(F,size=9,italic=True,color="666666"); c.alignment=Alignment(indent=1)
    ws.row_dimensions[2].height=18

def enc(ws,fila,vals,anchos=None):
    for j,v in enumerate(vals,1):
        c=ws.cell(fila,j,v); c.font=Font(F,size=9,bold=True,color="FFFFFF"); c.fill=CAB
        c.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True); c.border=BOX
    ws.row_dimensions[fila].height=30
    if anchos:
        for j,a in enumerate(anchos,1): ws.column_dimensions[get_column_letter(j)].width=a

# ───────────────────────── 1 · SUPUESTOS ─────────────────────────
ws=wb.active; ws.title="Supuestos"
titulo(ws,"QUORELIA · MODELO DE PRECIOS — CUENTA CATL VÍA CAVEX",
       "Editar SOLO las celdas amarillas. El azul es un dato de entrada; el negro se calcula solo.")
r=4
ws.cell(r,1,"LEYENDA").font=Font(F,size=9,bold=True)
for j,(t,f,col) in enumerate([("Celda a editar",AM,NG),("Dato de entrada",None,AZ),("Calculado",None,NG),("De otra hoja",None,VD)]):
    c=ws.cell(r,2+j,t); c.font=Font(F,size=8,color=col); c.border=BOX
    if f: c.fill=f
    ws.column_dimensions[get_column_letter(2+j)].width=17
ws.column_dimensions["A"].width=52; ws.column_dimensions["F"].width=46
r=6
def bloque(r,tit):
    ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=5)
    c=ws.cell(r,1,tit); c.font=Font(F,size=10,bold=True); c.fill=GRIS; c.alignment=Alignment(indent=1)
    ws.row_dimensions[r].height=20
    return r+1
def sup(r,etq,val,fmt,fuente="",edit=True):
    ws.cell(r,1,etq).font=Font(F,size=10)
    c=ws.cell(r,2,val); c.font=Font(F,size=10,bold=True,color=AZ); c.number_format=fmt; c.border=BOX
    if edit: c.fill=AM
    ws.cell(r,6,fuente).font=Font(F,size=8,italic=True,color="777777")
    return r+1

r=bloque(r,"MACRO — CHILE, AGOSTO 2026")
FX=r
r=sup(r,"Tipo de cambio (CLP por USD)",915,'#,##0',"Rango observado 909–928 en agosto 2026 (Wise)")
JOR=r
r=sup(r,"Jornada laboral máxima (horas/semana)",42,'#,##0',"Ley 21.561, vigente desde el 26 de abril de 2026")
IMM=r
r=sup(r,"Ingreso mínimo mensual (CLP)",553553,'#,##0',"Vigente desde mayo 2026")
r=sup(r,"IVA",0.19,PCT,"Los precios de este modelo son netos, sin IVA")
r+=1
r=bloque(r,"CARTERA — LO QUE HAY QUE CONFIRMAR CON CAVEX")
fila_sitios=r
r=sup(r,"Número de sitios en el alcance",12,'#,##0',"SUPUESTO. Cavex no ha confirmado cómo se reparten los ~6 GW")
r=sup(r,"Potencia media por sitio (MW)",250,'#,##0',"SUPUESTO. Cambiar cuando Cavex entregue el listado")
fila_pot=r-1
fila_gw=r
r=sup(r,"Planta actual bajo contrato (MW)",1000,'#,##0',"Contrato vigente Cavex–Quorelia")
r+=1
r=bloque(r,"PRECIO ACTUAL DEL CONTRATO VIGENTE")
fila_actual=r
r=sup(r,"Tarifa mensual actual (CLP, neta)",300000,'#,##0',"Lo que Quorelia factura hoy a Cavex por 1 GW")
r+=1
r=bloque(r,"COSTOS — ESCENARIO A · SOLO CAPA DE DATOS (puntos 1, 2 parcial y 5)")
fila_a0=r
for etq,val,fte in [("Guardia técnica 24/7 — 4 ingenieros en rotación",3400000,"Bono de disponibilidad + horas efectivas de intervención"),
                    ("Analista de operaciones (1 FTE cargado)",2200000,"Revisa hallazgos, emite informes, atiende a Cavex"),
                    ("Jefatura técnica / QA (0,5 FTE)",1500000,"Control de calidad de la cifra publicada"),
                    ("Infraestructura (hosting, ingesta, almacenamiento)",1200000,"Redundancia incluida"),
                    ("Licencias y herramientas (alerting, ticketing)",600000,""),
                    ("Seguro de responsabilidad civil profesional (E&O)",800000,"NUEVO — obligatorio si se asume alcance 24/7")]:
    r=sup(r,etq,val,'#,##0',fte)
fila_a1=r-1
ws.cell(r,1,"Overhead administrativo y comercial").font=Font(F,size=10)
c=ws.cell(r,2,0.20); c.font=Font(F,size=10,bold=True,color=AZ); c.number_format=PCT; c.fill=AM; c.border=BOX
fila_oh=r; r+=1
ws.cell(r,1,"COSTO FIJO MENSUAL — ESCENARIO A").font=Font(F,size=10,bold=True)
c=ws.cell(r,2,f"=SUM(B{fila_a0}:B{fila_a1})*(1+B{fila_oh})"); c.font=Font(F,size=10,bold=True); c.number_format=CLP; c.fill=TEAL; c.border=BOX
fila_costoA=r; r+=1
r=sup(r,"Costo marginal por sitio adicional",200000,'#,##0',"Ingesta, almacenamiento y tiempo de analista por hallazgo")
fila_margA=r-1
r+=1
r=bloque(r,"COSTOS — ESCENARIO B · CENTRO 24/7 CON DOTACIÓN (los 5 puntos del correo)")
r=sup(r,"Puestos concurrentes 24/7",2,'#,##0',"Uno solo no cubre alarmas + llamadas + órdenes a la vez")
fila_puestos=r-1
r=sup(r,"Holgura por vacaciones, licencias y capacitación",0.15,PCT,"15 días hábiles de vacaciones + licencias + formación")
fila_holg=r-1
ws.cell(r,1,"FTE necesarios por puesto 24/7").font=Font(F,size=10)
c=ws.cell(r,2,f"=168/B{JOR}*(1+B{fila_holg})"); c.font=Font(F,size=10,bold=True); c.number_format=NUM; c.border=BOX
ws.cell(r,6,"168 horas semanales ÷ jornada legal, más la holgura").font=Font(F,size=8,italic=True,color="777777")
fila_fte=r; r+=1
r=sup(r,"Costo cargado por operador de turno (CLP/mes)",1600000,'#,##0',"Bruto + turnos rotativos nocturnos + cargas del empleador")
fila_op=r-1
r=sup(r,"Ingenieros senior / jefes de turno",3,'#,##0',"")
fila_srn=r-1
r=sup(r,"Costo cargado por senior (CLP/mes)",2600000,'#,##0',"")
fila_src=r-1
r=sup(r,"Telefonía y atención multilingüe (ES/PT/EN/ZH)",1500000,'#,##0',"Punto 4 del correo: atención de llamadas 24/7")
fila_tel=r-1
r=sup(r,"Infraestructura y licencias del centro",2100000,'#,##0',"")
fila_infb=r-1
r=sup(r,"Seguro RC ampliado (juicio operacional sobre alarmas BMS)",1800000,'#,##0',"Punto 2: «preliminary judgments» sobre baterías de litio")
fila_segb=r-1
ws.cell(r,1,"COSTO FIJO MENSUAL — ESCENARIO B").font=Font(F,size=10,bold=True)
c=ws.cell(r,2,f"=(ROUND(B{fila_puestos}*B{fila_fte},0)*B{fila_op}+B{fila_srn}*B{fila_src}+B{fila_tel}+B{fila_infb}+B{fila_segb})*(1+B{fila_oh})")
c.font=Font(F,size=10,bold=True); c.number_format=CLP; c.fill=ROJO; c.border=BOX
fila_costoB=r; r+=1
r=sup(r,"Costo marginal por sitio adicional (escenario B)",300000,'#,##0',"")
fila_margB=r-1

# ───────────────────────── 2 · DÓNDE ESTÁ EL PRECIO HOY ─────────────────────────
FXR=FX  # fila del tipo de cambio en Supuestos
ws2=wb.create_sheet("Precio actual vs mercado")
titulo(ws2,"DÓNDE ESTÁ EL PRECIO ACTUAL","Normalización del contrato vigente y comparación contra la tabla de referencia del propio Quorelia.",6)
enc(ws2,4,["Métrica","Valor","Unidad","Cómo se calcula","",""],[46,17,15,52,3,3])
filas=[("Tarifa mensual actual",f"=Supuestos!B{fila_actual}","CLP/mes","Contrato vigente Cavex–Quorelia",CLP),
       ("Tarifa mensual actual en dólares",f"=Supuestos!B{fila_actual}/Supuestos!B{FXR}","USD/mes","÷ tipo de cambio",USD),
       ("Ingreso anual del contrato",f"=Supuestos!B{fila_actual}*12","CLP/año","× 12",CLP),
       ("Ingreso anual en dólares",f"=Supuestos!B{fila_actual}*12/Supuestos!B{FXR}","USD/año","",USD),
       ("Precio por MW",f"=Supuestos!B{fila_actual}/Supuestos!B{fila_gw}","CLP/MW/mes","÷ 1.000 MW contratados",'#,##0'),
       ("Precio por MW al año",f"=Supuestos!B{fila_actual}*12/Supuestos!B{fila_gw}/Supuestos!B{FXR}","USD/MW/año","",'"US$"#,##0.00'),
       ("Como fracción de UN sueldo mínimo",f"=Supuestos!B{fila_actual}/Supuestos!B{IMM}","× IMM","El ingreso mensual completo de esta cuenta",'0.00"×"')]
rr=5
for etq,fx,un,nota,fmt in filas:
    ws2.cell(rr,1,etq).font=Font(F,size=10)
    c=ws2.cell(rr,2,fx); c.font=Font(F,size=10,bold=True); c.number_format=fmt; c.border=BOX
    ws2.cell(rr,3,un).font=Font(F,size=9,color="777777")
    ws2.cell(rr,4,nota).font=Font(F,size=8,italic=True,color="777777")
    rr+=1
rr+=1
ws2.cell(rr,1,"CONTRA LA TABLA DE REFERENCIA QUE YA TENÍAS").font=Font(F,size=10,bold=True); rr+=1
enc(ws2,rr,["Proveedor","Piso indicativo (USD/año)","Quorelia hoy","Cuántas veces más caro es el piso","",""])
rr+=1
base=[("Quorelia · objetivo propio",20000),("AVEVA PI System",25000),("AlsoEnergy",30000),
      ("Power Factors / Unity",50000),("ABB",75000),("Siemens · Hitachi · GE Vernova",100000)]
prim=rr
for nom,piso in base:
    ws2.cell(rr,1,nom).font=Font(F,size=10,bold=(nom.startswith("Quorelia")))
    c=ws2.cell(rr,2,piso); c.font=Font(F,size=10,color=AZ); c.number_format=USD; c.border=BOX; c.fill=AM
    c=ws2.cell(rr,3,f"=Supuestos!B{fila_actual}*12/Supuestos!B{FXR}/B{rr}"); c.font=Font(F,size=10); c.number_format=PCT; c.border=BOX
    c=ws2.cell(rr,4,f"=B{rr}/(Supuestos!B{fila_actual}*12/Supuestos!B{FXR})"); c.font=Font(F,size=10,bold=True); c.number_format='0.0"×"'; c.border=BOX; c.fill=TEAL
    rr+=1
rr+=1
ws2.cell(rr,1,"LO QUE ESTABAS PENSANDO COBRAR").font=Font(F,size=10,bold=True); rr+=1
enc(ws2,rr,["Tarifa propuesta (CLP/sitio/mes)","USD/sitio/año","% de tu propio piso de US$20.000","","",""])
rr+=1
for p in (500000,1000000):
    c=ws2.cell(rr,1,p); c.font=Font(F,size=10,color=AZ); c.number_format=CLP; c.fill=AM; c.border=BOX
    c=ws2.cell(rr,2,f"=A{rr}*12/Supuestos!B{FXR}"); c.font=Font(F,size=10,bold=True); c.number_format=USD; c.border=BOX
    c=ws2.cell(rr,3,f"=B{rr}/B{prim}"); c.font=Font(F,size=10); c.number_format=PCT; c.border=BOX; c.fill=ROJO
    rr+=1
rr+=1
ws2.merge_cells(start_row=rr,start_column=1,end_row=rr,end_column=4)
c=ws2.cell(rr,1,"Incluso el techo de tu rango (1.000.000 CLP/sitio/mes) queda por debajo del piso de tu propia tabla de referencia.")
c.font=Font(F,size=9,bold=True,color="B00000"); c.alignment=Alignment(indent=1)

# ───────────────────────── 3 · TARIFA Y ESCENARIOS ─────────────────────────
ws3=wb.create_sheet("Tarifa recomendada")
titulo(ws3,"TARIFA RECOMENDADA — TRES BLOQUES","Separar la plataforma (escala) del servicio 24/7 (no escala). Amarillo = editable.",7)
r3=4
ws3.column_dimensions["A"].width=44
for L,w in zip("BCDEFG",[17,15,17,17,15,15]): ws3.column_dimensions[L].width=w
ws3.merge_cells(start_row=r3,start_column=1,end_row=r3,end_column=7)
c=ws3.cell(r3,1,"BLOQUE 1 · Plataforma, verificación y reportería — recurrente por sitio"); c.font=Font(F,size=10,bold=True); c.fill=GRIS; c.alignment=Alignment(indent=1); r3+=1
enc(ws3,r3,["Banda de potencia","CLP/sitio/mes","USD/sitio/mes","USD/sitio/año","","",""]); r3+=1
b1_0=r3
for banda,precio in [("Hasta 200 MW",950000),("200 – 500 MW",1450000),("Más de 500 MW",1950000)]:
    ws3.cell(r3,1,banda).font=Font(F,size=10)
    c=ws3.cell(r3,2,precio); c.font=Font(F,size=10,bold=True,color=AZ); c.number_format=CLP; c.fill=AM; c.border=BOX
    c=ws3.cell(r3,3,f"=B{r3}/Supuestos!B{FXR}"); c.font=Font(F,size=10); c.number_format=USD; c.border=BOX
    c=ws3.cell(r3,4,f"=B{r3}*12/Supuestos!B{FXR}"); c.font=Font(F,size=10,bold=True); c.number_format=USD; c.border=BOX
    r3+=1
b1_media=r3-2
r3+=1
ws3.merge_cells(start_row=r3,start_column=1,end_row=r3,end_column=7)
c=ws3.cell(r3,1,"BLOQUE 2 · Servicio de monitoreo y guardia 24/7 — cargo fijo de cuenta"); c.font=Font(F,size=10,bold=True); c.fill=GRIS; c.alignment=Alignment(indent=1); r3+=1
ws3.cell(r3,1,"Cargo fijo mensual (hasta 20 sitios)").font=Font(F,size=10)
c=ws3.cell(r3,2,14500000); c.font=Font(F,size=10,bold=True,color=AZ); c.number_format=CLP; c.fill=AM; c.border=BOX
c=ws3.cell(r3,3,f"=B{r3}/Supuestos!B{FXR}"); c.font=Font(F,size=10); c.number_format=USD; c.border=BOX
fijo=r3; r3+=1
ws3.cell(r3,1,"Incremento por sitio a partir del 21º").font=Font(F,size=10)
c=ws3.cell(r3,2,350000); c.font=Font(F,size=10,bold=True,color=AZ); c.number_format=CLP; c.fill=AM; c.border=BOX
extra=r3; r3+=2
ws3.merge_cells(start_row=r3,start_column=1,end_row=r3,end_column=7)
c=ws3.cell(r3,1,"BLOQUE 3 · Implementación — pago único por sitio"); c.font=Font(F,size=10,bold=True); c.fill=GRIS; c.alignment=Alignment(indent=1); r3+=1
enc(ws3,r3,["Tipo de integración","CLP (una vez)","USD","","","",""]); r3+=1
impl0=r3
for t,v in [("Estándar — SCADA/EMS conocido",6500000),("Compleja — BMS propietario o multi-vendor",13000000)]:
    ws3.cell(r3,1,t).font=Font(F,size=10)
    c=ws3.cell(r3,2,v); c.font=Font(F,size=10,bold=True,color=AZ); c.number_format=CLP; c.fill=AM; c.border=BOX
    c=ws3.cell(r3,3,f"=B{r3}/Supuestos!B{FXR}"); c.font=Font(F,size=10); c.number_format=USD; c.border=BOX
    r3+=1
r3+=1
ws3.merge_cells(start_row=r3,start_column=1,end_row=r3,end_column=7)
c=ws3.cell(r3,1,"CÓMO QUEDA SEGÚN EL TAMAÑO DE LA CARTERA — escenario A (solo capa de datos)")
c.font=Font(F,size=10,bold=True); c.fill=GRIS; c.alignment=Alignment(indent=1); r3+=1
enc(ws3,r3,["Sitios en cartera","Recurrente CLP/mes","USD/año","USD/sitio/año","Costo CLP/mes","Resultado CLP/mes","Margen bruto"]); r3+=1
for n in (1,6,12,20,30,40):
    banda_ref=b1_0+2 if n==1 else (b1_media if n<=20 else b1_0)
    ws3.cell(r3,1,n).font=Font(F,size=10,bold=True)
    c=ws3.cell(r3,2,f"=A{r3}*B{banda_ref}+B{fijo}+MAX(0,A{r3}-20)*B{extra}"); c.font=Font(F,size=10); c.number_format=CLP; c.border=BOX
    c=ws3.cell(r3,3,f"=B{r3}*12/Supuestos!B{FXR}"); c.font=Font(F,size=10); c.number_format=USD; c.border=BOX
    c=ws3.cell(r3,4,f"=C{r3}/A{r3}"); c.font=Font(F,size=10); c.number_format=USD; c.border=BOX
    c=ws3.cell(r3,5,f"=Supuestos!B{fila_costoA}+A{r3}*Supuestos!B{fila_margA}"); c.font=Font(F,size=10,color=VD); c.number_format=CLP; c.border=BOX
    c=ws3.cell(r3,6,f"=B{r3}-E{r3}"); c.font=Font(F,size=10,bold=True); c.number_format=CLP; c.border=BOX
    c=ws3.cell(r3,7,f"=IF(B{r3}=0,0,F{r3}/B{r3})"); c.font=Font(F,size=10,bold=True); c.number_format=PCT; c.border=BOX; c.fill=TEAL
    r3+=1

# ───────────────────────── 4 · COMPARACIÓN A vs B ─────────────────────────
ws4=wb.create_sheet("Escenario A vs B")
titulo(ws4,"QUÉ PASA SI ACEPTAS LOS CINCO PUNTOS DEL CORREO","Escenario A = capa de datos (1, 2 parcial, 5). Escenario B = centro 24/7 con dotación (los 5).",6)
ws4.column_dimensions["A"].width=50
for L in "BCDEF": ws4.column_dimensions[L].width=19
r4=4
enc(ws4,r4,["Punto del correo de CATL","Naturaleza","¿Quorelia o Cavex?","Por qué","",""]); r4+=1
alcance=[("1 · Monitoreo en tiempo real 24/7 de BMS, PCS y EMS","Software","QUORELIA","Es exactamente el producto. Costo marginal casi nulo."),
 ("2 · Manejo de alarmas: detectar, clasificar, notificar, registrar","Software","QUORELIA","La detección y el registro trazable son el diferenciador."),
 ("2b · «Preliminary judgments» sobre la alarma","Decisión operacional","CAVEX","Juicio sobre baterías de litio. Contradice tus Términos §03."),
 ("3 · Órdenes de trabajo: distribuir y gestionar 24/7","Coordinación de terreno","CAVEX","Trabajo humano de campo. Quorelia puede generar el hallazgo, no despachar cuadrillas."),
 ("4 · Atención de llamadas 24/7 a cliente y personal en sitio","Call center","CAVEX","Es un call center multilingüe. Margen bajo, nada que ver con tu producto."),
 ("5 · Registro de datos e informe diario de operación","Software","QUORELIA","Ya lo entregas hoy. Es el corazón del servicio.")]
for p,nat,quien,pq in alcance:
    ws4.cell(r4,1,p).font=Font(F,size=9); ws4.cell(r4,1).alignment=Alignment(wrap_text=True,vertical="top")
    ws4.cell(r4,2,nat).font=Font(F,size=9)
    c=ws4.cell(r4,3,quien); c.font=Font(F,size=9,bold=True); c.alignment=Alignment(horizontal="center")
    c.fill=TEAL if quien=="QUORELIA" else ROJO
    ws4.cell(r4,4,pq).font=Font(F,size=8,color="555555"); ws4.cell(r4,4).alignment=Alignment(wrap_text=True,vertical="top")
    ws4.row_dimensions[r4].height=30
    for j in range(1,5): ws4.cell(r4,j).border=BOX
    r4+=1
r4+=1
enc(ws4,r4,["Comparación","Escenario A · capa de datos","Escenario B · centro dotado","Diferencia","",""]); r4+=1
comp=[("Costo fijo mensual (CLP)",f"=Supuestos!B{fila_costoA}",f"=Supuestos!B{fila_costoB}",CLP),
      ("Costo fijo anual (CLP)",f"=B{r4}*12",f"=C{r4}*12",CLP),
      ("Costo fijo anual (USD)",f"=B{r4}*12/Supuestos!B{FXR}",f"=C{r4}*12/Supuestos!B{FXR}",USD),
      ("Personas en nómina dedicadas","=5.5",f"=ROUND(Supuestos!B{fila_puestos}*Supuestos!B{fila_fte},0)+Supuestos!B{fila_srn}",NUM),
      ("Sitios para equilibrio a 1.000.000 CLP/sitio/mes",f"=ROUNDUP(Supuestos!B{fila_costoA}/(1000000-Supuestos!B{fila_margA}),0)",f"=ROUNDUP(Supuestos!B{fila_costoB}/(1000000-Supuestos!B{fila_margB}),0)",'#,##0'),
      ("Sitios para equilibrio a 500.000 CLP/sitio/mes",f"=ROUNDUP(Supuestos!B{fila_costoA}/(500000-Supuestos!B{fila_margA}),0)",f"=ROUNDUP(Supuestos!B{fila_costoB}/(500000-Supuestos!B{fila_margB}),0)",'#,##0')]
base4=r4
for etq,fa,fb,fmt in comp:
    ws4.cell(r4,1,etq).font=Font(F,size=10)
    c=ws4.cell(r4,2,fa); c.font=Font(F,size=10,bold=True); c.number_format=fmt; c.border=BOX; c.fill=TEAL
    c=ws4.cell(r4,3,fb); c.font=Font(F,size=10,bold=True); c.number_format=fmt; c.border=BOX; c.fill=ROJO
    c=ws4.cell(r4,4,f"=C{r4}-B{r4}"); c.font=Font(F,size=10); c.number_format=fmt; c.border=BOX
    r4+=1
r4+=1
ws4.merge_cells(start_row=r4,start_column=1,end_row=r4,end_column=4)
c=ws4.cell(r4,1,"El escenario B no es una ampliación del servicio: es otro negocio, con costo laboral que no escala y margen de BPO.")
c.font=Font(F,size=9,bold=True,color="B00000"); c.alignment=Alignment(indent=1)
r4+=2
ws4.merge_cells(start_row=r4,start_column=1,end_row=r4,end_column=4)
c=ws4.cell(r4,1,"QUÉ PASA CON LO QUE IBAS A COTIZAR — cartera de 12 sitios"); c.font=Font(F,size=10,bold=True); c.fill=GRIS; c.alignment=Alignment(indent=1); r4+=1
enc(ws4,r4,["Opción de precio","Ingreso CLP/mes","Costo CLP/mes","Resultado CLP/mes","Margen",""]); r4+=1
opc=[("Tu rango bajo — 500.000 CLP/sitio","=12*500000"),("Tu rango alto — 1.000.000 CLP/sitio","=12*1000000"),
     ("Tarifa recomendada (bloques 1+2)",f"=12*'Tarifa recomendada'!B{b1_media}+'Tarifa recomendada'!B{fijo}")]
for etq,fx in opc:
    ws4.cell(r4,1,etq).font=Font(F,size=10)
    c=ws4.cell(r4,2,fx); c.font=Font(F,size=10,bold=True); c.number_format=CLP; c.border=BOX
    c=ws4.cell(r4,3,f"=Supuestos!B{fila_costoA}+12*Supuestos!B{fila_margA}"); c.font=Font(F,size=10,color=VD); c.number_format=CLP; c.border=BOX
    c=ws4.cell(r4,4,f"=B{r4}-C{r4}"); c.font=Font(F,size=10,bold=True); c.number_format=CLP; c.border=BOX
    c=ws4.cell(r4,5,f"=D{r4}/B{r4}"); c.font=Font(F,size=10,bold=True); c.number_format=PCT; c.border=BOX
    r4+=1

for s in wb.worksheets: s.sheet_view.showGridLines=False
wb.save("Modelo-Precios-Quorelia-CATL.xlsx")
print("escrito")
